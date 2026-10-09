#!/usr/bin/env python3
"""Inventaria e extrai localmente os documentos de referência.

O script não altera os originais. Transcrições integrais ficam em private/,
fora do Git e do build público.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import sys
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import openpyxl
import pdfplumber
from docx import Document


APP_DIR = Path(__file__).resolve().parents[1]
WORKSPACE_DIR = APP_DIR.parent
SOURCE_DIR = WORKSPACE_DIR / "arquivos"
PRIVATE_DIR = APP_DIR / "private" / "analise-documental"
DOCS_DIR = APP_DIR / "docs-tecnicos"
INVENTORY_JSON = DOCS_DIR / "inventario-documental.json"
INVENTORY_MD = DOCS_DIR / "INVENTARIO_DOCUMENTAL.md"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def stable_slug(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-zA-Z0-9]+", "-", ascii_value).strip("-").lower()


def write_text_atomic(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(content, encoding="utf-8", newline="\n")
    temporary.replace(path)


def write_json_atomic(path: Path, payload: Any) -> None:
    write_text_atomic(path, json.dumps(payload, ensure_ascii=False, indent=2) + "\n")


def extract_pdf(path: Path) -> dict[str, Any]:
    document_dir = PRIVATE_DIR / stable_slug(path.stem)
    document_dir.mkdir(parents=True, exist_ok=True)
    pages: list[dict[str, Any]] = []
    combined: list[str] = []
    with pdfplumber.open(path) as pdf:
        metadata = {str(k): str(v) for k, v in (pdf.metadata or {}).items()}
        for number, page in enumerate(pdf.pages, start=1):
            error = None
            try:
                text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            except Exception as exc:  # uma falha deve ficar explícita no inventário
                text = ""
                error = f"{type(exc).__name__}: {exc}"
            page_path = document_dir / f"pagina-{number:03d}.txt"
            write_text_atomic(page_path, text + ("\n" if text else ""))
            ocr_status = "pendente" if len(text.strip()) < 150 else "nao_necessario"
            ocr_details: dict[str, Any] | None = None
            if ocr_status == "pendente" and os.environ.get("ENABLE_LOCAL_OCR") == "1":
                try:
                    import fitz
                    from rapidocr_onnxruntime import RapidOCR

                    pdf_document = fitz.open(path)
                    pixmap = pdf_document[number - 1].get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
                    image_path = document_dir / f"pagina-{number:03d}-ocr-fonte.png"
                    pixmap.save(image_path)
                    pdf_document.close()
                    engine = RapidOCR()
                    result, elapsed = engine(str(image_path))
                    ocr_lines = [] if not result else [str(line[1]) for line in result if len(line) >= 2]
                    ocr_text = "\n".join(ocr_lines)
                    ocr_path = document_dir / f"pagina-{number:03d}-ocr.txt"
                    write_text_atomic(ocr_path, ocr_text + ("\n" if ocr_text else ""))
                    ocr_status = "executado_com_texto" if ocr_text.strip() else "executado_sem_texto"
                    ocr_details = {
                        "caracteres": len(ocr_text),
                        "duracao_segundos": elapsed,
                        "arquivo_texto": ocr_path.relative_to(APP_DIR).as_posix(),
                        "arquivo_imagem_fonte": image_path.relative_to(APP_DIR).as_posix(),
                    }
                except Exception as exc:
                    ocr_status = "erro"
                    ocr_details = {"erro": f"{type(exc).__name__}: {exc}"}
            pages.append(
                {
                    "pagina": number,
                    "caracteres_extraidos": len(text),
                    "linhas_extraidas": len(text.splitlines()),
                    "extracao_textual": "erro" if error else ("vazia" if not text.strip() else "extraida"),
                    "ocr_local": ocr_status,
                    "detalhes_ocr": ocr_details,
                    "erro": error,
                    "arquivo_privado": page_path.relative_to(APP_DIR).as_posix(),
                }
            )
            combined.append(f"===== PÁGINA {number} =====\n{text}\n")
    combined_path = document_dir / "texto-integral-por-pagina.txt"
    write_text_atomic(combined_path, "\n".join(combined))
    report = {
        "arquivo_origem": f"../arquivos/{path.name}",
        "paginas_declaradas": len(pages),
        "paginas_processadas": len(pages),
        "paginas_sem_texto": [item["pagina"] for item in pages if item["extracao_textual"] == "vazia"],
        "paginas_com_texto_insuficiente_para_descartar_ocr": [item["pagina"] for item in pages if item["ocr_local"] == "pendente"],
        "paginas_com_ocr_executado": [item["pagina"] for item in pages if item["ocr_local"].startswith("executado")],
        "paginas_com_erro_de_ocr": [item["pagina"] for item in pages if item["ocr_local"] == "erro"],
        "metadata_pdf": metadata,
        "paginas": pages,
    }
    write_json_atomic(document_dir / "relatorio-extracao.json", report)
    return {
        "quantidade_paginas": len(pages),
        "paginas_processadas": len(pages),
        "paginas_sem_texto": report["paginas_sem_texto"],
        "paginas_ocr_pendente": report["paginas_com_texto_insuficiente_para_descartar_ocr"],
        "caracteres_extraidos": sum(item["caracteres_extraidos"] for item in pages),
        "status_leitura": (
            "extraido_com_ocr_pendente"
            if report["paginas_com_texto_insuficiente_para_descartar_ocr"]
            else "extraido_com_erro_de_ocr"
            if report["paginas_com_erro_de_ocr"]
            else "texto_e_ocr_local_extraidos"
            if report["paginas_com_ocr_executado"]
            else "texto_extraido_integralmente"
        ),
        "relatorio_privado": (document_dir / "relatorio-extracao.json").relative_to(APP_DIR).as_posix(),
    }


def _docx_extended_properties(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    with zipfile.ZipFile(path) as archive:
        if "docProps/app.xml" not in archive.namelist():
            return result
        root = ET.fromstring(archive.read("docProps/app.xml"))
        for child in root:
            result[child.tag.rsplit("}", 1)[-1]] = child.text or ""
    return result


def _docx_xml_text(path: Path, member: str) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        if member not in archive.namelist():
            return []
        root = ET.fromstring(archive.read(member))
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    lines: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace))
        if text.strip():
            lines.append(text)
    return lines


def extract_docx(path: Path) -> dict[str, Any]:
    doc = Document(path)
    properties = _docx_extended_properties(path)
    paragraphs = [
        {"indice": index, "estilo": paragraph.style.name if paragraph.style else None, "texto": paragraph.text}
        for index, paragraph in enumerate(doc.paragraphs, start=1)
    ]
    tables: list[dict[str, Any]] = []
    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            rows.append(
                {
                    "linha": row_index,
                    "celulas": [cell.text for cell in row.cells],
                }
            )
        tables.append({"tabela": table_index, "linhas": len(rows), "colunas": len(table.columns), "conteudo": rows})
    headers: list[dict[str, Any]] = []
    footers: list[dict[str, Any]] = []
    for index, section in enumerate(doc.sections, start=1):
        headers.append({"secao": index, "texto": [p.text for p in section.header.paragraphs]})
        footers.append({"secao": index, "texto": [p.text for p in section.footer.paragraphs]})
    footnotes = _docx_xml_text(path, "word/footnotes.xml")
    endnotes = _docx_xml_text(path, "word/endnotes.xml")
    comments = _docx_xml_text(path, "word/comments.xml")
    with zipfile.ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")]
    payload = {
        "arquivo_origem": f"../arquivos/{path.name}",
        "propriedades_estendidas": properties,
        "paragrafos": paragraphs,
        "tabelas": tables,
        "cabecalhos": headers,
        "rodapes": footers,
        "notas_de_rodape": footnotes,
        "notas_finais": endnotes,
        "comentarios": comments,
        "midias_incorporadas": media,
    }
    output = PRIVATE_DIR / "roteiro-docx-extraido.json"
    write_json_atomic(output, payload)
    page_count = int(properties.get("Pages", "0") or 0) or None
    return {
        "quantidade_paginas": page_count,
        "quantidade_paragrafos": len(paragraphs),
        "quantidade_tabelas": len(tables),
        "quantidade_secoes": len(doc.sections),
        "midias_incorporadas": len(media),
        "status_leitura": "estrutura_extraida_integralmente",
        "relatorio_privado": output.relative_to(APP_DIR).as_posix(),
    }


def excel_serializable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def extract_xlsx(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    values_workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        values_sheet = values_workbook[worksheet.title]
        cells: list[dict[str, Any]] = []
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                cells.append(
                    {
                        "coordenada": cell.coordinate,
                        "linha": cell.row,
                        "coluna": cell.column,
                        "valor_ou_formula": excel_serializable(cell.value),
                        "valor_calculado_em_cache": excel_serializable(values_sheet[cell.coordinate].value),
                        "tipo": cell.data_type,
                        "formato_numerico": cell.number_format,
                        "estilo_id": cell.style_id,
                    }
                )
        sheets.append(
            {
                "nome": worksheet.title,
                "linhas": worksheet.max_row,
                "colunas": worksheet.max_column,
                "dimensao": worksheet.calculate_dimension(),
                "celulas_mescladas": [str(item) for item in worksheet.merged_cells.ranges],
                "celulas": cells,
            }
        )
    payload = {
        "arquivo_origem": f"../arquivos/{path.name}",
        "abas": sheets,
        "nomes_definidos": [str(item) for item in workbook.defined_names.values()],
    }
    output = PRIVATE_DIR / "planilha-xlsx-extraida.json"
    write_json_atomic(output, payload)
    return {
        "quantidade_abas": len(sheets),
        "abas": [{"nome": item["nome"], "linhas": item["linhas"], "colunas": item["colunas"]} for item in sheets],
        "status_leitura": "celulas_formulas_estilos_extraidos_integralmente",
        "relatorio_privado": output.relative_to(APP_DIR).as_posix(),
    }


def build_markdown(inventory: dict[str, Any]) -> str:
    lines = [
        "# Inventário documental",
        "",
        "> Inventário técnico gerado localmente. Os originais permanecem em `../arquivos/` e não são publicados.",
        "",
        f"- Arquivos encontrados: **{inventory['totais']['arquivos']}**",
        f"- PDFs: **{inventory['totais']['pdfs']}**",
        f"- Planilhas XLSX: **{inventory['totais']['xlsx']}**",
        f"- Documentos DOCX: **{inventory['totais']['docx']}**",
        f"- Páginas PDF processadas: **{inventory['totais']['paginas_pdf']}**",
        "",
        "| Arquivo original | Tipo | Tamanho (bytes) | Páginas/abas | SHA-256 | Status |",
        "|---|---:|---:|---:|---|---|",
    ]
    for item in inventory["arquivos"]:
        count = item.get("quantidade_paginas") or item.get("quantidade_abas") or "—"
        lines.append(
            f"| {item['nome'].replace('|', '&#124;')} | {item['extensao']} | {item['tamanho_bytes']} | {count} | `{item['sha256']}` | {item['status_leitura']} |"
        )
    lines.extend(
        [
            "",
            "## Limitações registradas",
            "",
            "- A paginação do DOCX é a informada nas propriedades do próprio arquivo e pode variar conforme o motor de renderização.",
            "- Páginas com menos de 150 caracteres extraídos são marcadas para inspeção/OCR local complementar; isso não autoriza inferir conteúdo ausente.",
            "- Os extratos integrais são restritos e ficam em `private/analise-documental/`, excluídos do Git e do build.",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> int:
    if not SOURCE_DIR.is_dir():
        print(f"Pasta de origem não encontrada: {SOURCE_DIR}", file=sys.stderr)
        return 2
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    PRIVATE_DIR.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    for path in sorted((item for item in SOURCE_DIR.iterdir() if item.is_file()), key=lambda item: item.name.casefold()):
        record: dict[str, Any] = {
            "nome": path.name,
            "caminho_relativo": f"../arquivos/{path.name}",
            "extensao": path.suffix.lower().lstrip("."),
            "tamanho_bytes": path.stat().st_size,
            "sha256": sha256_file(path),
        }
        try:
            if path.suffix.lower() == ".pdf":
                record.update(extract_pdf(path))
            elif path.suffix.lower() == ".docx":
                record.update(extract_docx(path))
            elif path.suffix.lower() == ".xlsx":
                record.update(extract_xlsx(path))
            else:
                record["status_leitura"] = "formato_nao_processado"
        except Exception as exc:
            record.update({"status_leitura": "erro", "erro": f"{type(exc).__name__}: {exc}"})
        records.append(record)
        print(f"{record['status_leitura']}: {path.name}")
    inventory = {
        "schema_version": 1,
        "gerado_em_utc": datetime.now(timezone.utc).isoformat(),
        "pasta_origem_relativa": "../arquivos",
        "totais": {
            "arquivos": len(records),
            "pdfs": sum(item["extensao"] == "pdf" for item in records),
            "xlsx": sum(item["extensao"] == "xlsx" for item in records),
            "docx": sum(item["extensao"] == "docx" for item in records),
            "paginas_pdf": sum(item.get("quantidade_paginas", 0) or 0 for item in records if item["extensao"] == "pdf"),
        },
        "arquivos": records,
    }
    write_json_atomic(INVENTORY_JSON, inventory)
    write_text_atomic(INVENTORY_MD, build_markdown(inventory))
    return 1 if any(item["status_leitura"] == "erro" for item in records) else 0


if __name__ == "__main__":
    raise SystemExit(main())
