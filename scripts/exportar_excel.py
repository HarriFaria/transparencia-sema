#!/usr/bin/env python3
"""Preenche uma copia do XLSX original com revisoes humanas consolidadas."""

from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


APP = Path(__file__).resolve().parents[1]
SOURCE = next((APP.parent / "arquivos").glob("*.xlsx"))
CONSOLIDATED = APP / "data" / "validacoes" / "consolidado.json"
EXPORTS = APP / "exports"


def latest_reviews() -> dict[str, dict]:
    if not CONSOLIDATED.exists():
        return {}
    data = json.loads(CONSOLIDATED.read_text(encoding="utf-8"))
    latest: dict[str, dict] = {}
    for review in data.get("revisions", []):
        if review.get("estado") != "validado" or not review.get("validado_pelo_auditor"):
            continue
        current = latest.get(review["criterio_id"])
        if current is None or review["registrado_em"] > current["registrado_em"]:
            latest[review["criterio_id"]] = review
    return latest


def main() -> None:
    reviews = latest_reviews()
    workbook = openpyxl.load_workbook(SOURCE, data_only=False)
    sheet = workbook.worksheets[0]
    by_id = {sheet.cell(row, 3).value: row for row in range(2, 75)}
    for criterion_id, review in reviews.items():
        if criterion_id not in by_id:
            raise RuntimeError(f"Criterio nao localizado no modelo: {criterion_id}")
        row = by_id[criterion_id]
        instant = datetime.fromisoformat(review["registrado_em"].replace("Z", "+00:00"))
        values = {
            13: instant.date(),
            14: instant.time().replace(microsecond=0),
            15: review.get("diagnostico_id") or review.get("execucao_id"),
            17: "Sim" if review.get("evidencias_disponiveis") else "Nao",
            18: review["parecer"],
            19: review["justificativa"],
        }
        for column, value in values.items():
            sheet.cell(row, column).value = value
    EXPORTS.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    destination = EXPORTS / f"Lista_Verificacao_SEMA_exportada_{timestamp}.xlsx"
    workbook.save(destination)
    csv_destination = EXPORTS / f"Lista_Verificacao_SEMA_{timestamp}.csv"
    with csv_destination.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        for row in sheet.iter_rows(min_row=1, max_row=74, min_col=1, max_col=20, values_only=True):
            writer.writerow(["" if value is None else value for value in row])
    print(f"Exportado: {destination} e {csv_destination} ({len(reviews)} revisao(oes)); sem identificacao pessoal.")


if __name__ == "__main__":
    main()
